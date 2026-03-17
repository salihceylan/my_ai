import os
import httpx
from bs4 import BeautifulSoup


async def load_file(file_path: str) -> list[str]:
    """PDF, Word, Excel, TXT gibi dosyaları okur ve metin listesi döner."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext in [".txt", ".md", ".csv"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return [f.read()]

    elif ext == ".pdf":
        import pdfplumber
        texts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    texts.append(text)
        return texts

    elif ext in [".docx"]:
        from docx import Document
        doc = Document(file_path)
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        return texts

    elif ext in [".xlsx", ".xls"]:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        texts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(str(c) for c in row if c is not None)
                if row_text.strip():
                    texts.append(row_text)
        return texts

    else:
        raise ValueError(f"Desteklenmeyen format: {ext}")


async def load_url(url: str) -> list[str]:
    """Web sayfasını okur ve metin listesi döner."""
    async with httpx.AsyncClient(timeout=30) as client:
        response = await client.get(url, follow_redirects=True)
        response.raise_for_status()
        html = response.text

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    texts = []
    for elem in soup.find_all(["p", "h1", "h2", "h3", "h4", "li", "td"]):
        text = elem.get_text(strip=True)
        if len(text) > 30:
            texts.append(text)

    return texts
